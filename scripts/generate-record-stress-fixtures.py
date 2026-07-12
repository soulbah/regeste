from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


OUT = Path(__file__).resolve().parents[1] / "static/dev/record-stress"
OUT.mkdir(parents=True, exist_ok=True)

TRANSFERS = [
    ("TX-ALPHA-001", "20 juin 2026", "120,00 EUR", "1 200 000 GNF", "2,00 EUR", "122,00 EUR"),
    ("TX-BRAVO-002", "12 juin 2026", "80,00 EUR", "800 000 GNF", "2,00 EUR", "82,00 EUR"),
    ("TX-CHARLIE-003", "08 juin 2026", "120,00 EUR", "1 212 000 GNF", "2,00 EUR", "122,00 EUR"),
    ("TX-DELTA-004", "13 mai 2026", "400,00 EUR", "4 120 000 GNF", "3,00 EUR", "403,00 EUR"),
    ("TX-ECHO-005", "09 mai 2026", "90,00 EUR", "927 000 GNF", "1,00 EUR", "91,00 EUR"),
    ("TX-FOXTROT-006", "02 mai 2026", "310,00 EUR", "3 193 000 GNF", "2,00 EUR", "312,00 EUR"),
    ("TX-GOLF-007", "20 avril 2026", "250,00 EUR", "2 600 000 GNF", "0,00 EUR", "250,00 EUR"),
]


def build_pdf() -> None:
    path = OUT / "repeated-transfers.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    for index, (record_id, date, sent, received, fee, total) in enumerate(TRANSFERS, 1):
        pdf.setFont("Helvetica-Bold", 17)
        pdf.drawString(54, height - 62, "Recepisse de transfert")
        pdf.setFont("Helvetica", 9)
        pdf.drawRightString(width - 54, height - 58, f"PAGE {index} / 7")
        pdf.setStrokeColorRGB(0.75, 0.75, 0.75)
        pdf.line(54, height - 76, width - 54, height - 76)
        y = height - 112
        rows = [
            ("Expediteur", "Camille Exemple"),
            ("Transaction", f"N deg {record_id} du {date} a 10:30"),
            ("Montant", sent),
            ("Montant recu par le beneficiaire", received),
            ("Frais de transfert", fee),
            ("Total", total),
            ("Beneficiaire", "Morgan Exemple"),
            ("Motif", "Assistance familiale"),
        ]
        for label, value in rows:
            pdf.setFont("Helvetica-Bold", 10)
            pdf.drawString(64, y, label)
            pdf.setFont("Helvetica", 10)
            pdf.drawString(270, y, value)
            pdf.setStrokeColorRGB(0.9, 0.9, 0.9)
            pdf.line(64, y - 7, width - 64, y - 7)
            y -= 34
        pdf.setFont("Helvetica", 8)
        boilerplate = "Conditions generales identiques sur chaque recepisse. Document synthetique de test, sans donnees reelles."
        pdf.drawString(64, 92, boilerplate)
        pdf.drawString(64, 76, "Aucun conseil financier. Valeurs creees uniquement pour tester indexation et recherche.")
        pdf.showPage()
    pdf.save()


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def build_docx() -> None:
    path = OUT / "repeated-transfers.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"),
        ("Heading 3", 12, 10, 5, "1F4D78"),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph("Registre synthetique de transferts", style="Heading 1")
    title.paragraph_format.keep_with_next = True
    note = doc.add_paragraph("Fixture anonymisee. Sept sections quasi identiques testent extraction, indexation et calcul exact.")
    note.paragraph_format.keep_with_next = True

    for index, (record_id, date, sent, received, fee, total) in enumerate(TRANSFERS, 1):
        if index > 1:
            doc.add_page_break()
        heading = doc.add_paragraph(f"Transaction {record_id}", style="Heading 2")
        heading.paragraph_format.keep_with_next = True
        rows = [
            ("Reference", record_id),
            ("Date", date),
            ("Expediteur", "Camille Exemple"),
            ("Montant", sent),
            ("Montant recu", received),
            ("Frais de transfert", fee),
            ("Total", total),
            ("Beneficiaire", "Morgan Exemple"),
            ("Motif", "Assistance familiale"),
        ]
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        set_table_geometry(table, [2700, 6660])
        for row, (label, value) in zip(table.rows, rows):
            row.cells[0].text = label
            row.cells[1].text = value
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                if cell_index == 0:
                    paragraph.runs[0].bold = True
        boilerplate = doc.add_paragraph(
            "Conditions generales identiques sur chaque recu. Valeurs synthetiques, sans donnees personnelles."
        )
        boilerplate.paragraph_format.space_before = Pt(10)
    doc.save(path)


def build_text() -> None:
    blocks = []
    for record_id, date, sent, received, fee, total in TRANSFERS:
        blocks.append(
            "\n".join(
                [
                    f"Transaction {record_id}",
                    f"Date: {date}",
                    "Expediteur: Camille Exemple",
                    f"Montant {sent} Montant recu par le beneficiaire {received}",
                    f"Frais de transfert {fee}",
                    f"Total {total}",
                    "Beneficiaire: Morgan Exemple",
                    "Motif: Assistance familiale",
                ]
            )
        )
    (OUT / "repeated-transfers.txt").write_text("\n\n".join(blocks) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_pdf()
    build_docx()
    build_text()
