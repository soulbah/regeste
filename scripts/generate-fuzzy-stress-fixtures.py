"""Generate controlled multi-format retrieval fixtures. Fictional truth only."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path("static/dev/fuzzy-stress")
BLUE = RGBColor(0x2E, 0x74, 0xB5)


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_width(cell, width: int) -> None:
    tcw = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        cell._tc.get_or_add_tcPr().append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def setup(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before, after in [("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6)]:
        style = styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), BLUE
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    heading = doc.add_heading(title, 0)
    heading.style = styles["Heading 1"]


def table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
    if width.getparent() is None:
        props.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    props.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for cell, value in zip(row.cells, widths):
            set_cell_width(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def profile_docx(path: Path, status: str, fact: str, revision: str) -> None:
    doc = Document()
    setup(doc, "Dossier fictif — Malik Ouedraogo")
    doc.add_heading("Identité", 1)
    doc.add_paragraph("Personne concernée : Malik Ouedraogo (alias M. Ouedraogo).")
    doc.add_heading("Situation déclarée", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["Révision", "Statut", "Preuve"]):
        cell.text = text
        shade(cell, "F2F4F7")
        cell.paragraphs[0].runs[0].bold = True
    row = table.add_row().cells
    for cell, text in zip(row, [revision, status, fact]):
        cell.text = text
    table_geometry(table, [1440, 2160, 5760])
    doc.add_heading("Limite", 1)
    doc.add_paragraph("Ce document ne contient aucune information sur sa profession ni son adresse.")
    doc.save(path)


def inventory_docx(path: Path) -> None:
    doc = Document()
    setup(doc, "Inventaire de caisses — jeu contrôlé")
    doc.add_paragraph("Les en-têtes définissent le sens de chaque valeur. Les références proches ne sont pas interchangeables.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    rows = [
        ["Référence", "Capacité", "Classe", "Commentaire"],
        ["BX-77", "42 kg", "renforcée", "usage intérieur"],
        ["BX-71", "24 kg", "standard", "ne pas confondre avec BX-77"],
        ["8X-77", "7 kg", "prototype", "OCR volontaire : huit, pas B"],
    ]
    for values in rows:
        cells = table.rows[0].cells if values is rows[0] else table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            if values is rows[0]:
                shade(cell, "E8EEF5")
                cell.paragraphs[0].runs[0].bold = True
    table_geometry(table, [1800, 1440, 1800, 4320])
    note = doc.add_paragraph("Note : une capacité absente ne doit jamais être déduite d'une référence similaire.")
    note.paragraph_format.space_before = Pt(8)
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    profile_docx(OUT / "malik-profile-study.docx", "Étudiant", "Inscription universitaire active pour 2026-2027.", "R1")
    profile_docx(OUT / "malik-profile-family.docx", "Marié", "Acte fictif enregistré le 14 mai 2025.", "R2")
    inventory_docx(OUT / "near-identifiers-table.docx")
    (OUT / "malik-profile-note.txt").write_text("Malik Ouedraogo pratique le violoncelle.\nAucune adresse n'est fournie.\n", encoding="utf-8")
    (OUT / "http-semantics.md").write_text("# Hypertext Transfer Protocol (HTTP)\n\nHTTP est un protocole sans état.\n\n## Négatif contrôlé\n\nAucun mot de passe n'est recommandé ici.\n", encoding="utf-8")


if __name__ == "__main__":
    main()
